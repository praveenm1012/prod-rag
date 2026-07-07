"""Shared fixtures for ingestion tests."""

from pathlib import Path

import pytest
from docx import Document as DocxDocument


@pytest.fixture
def sample_txt(tmp_path: Path) -> Path:
    path = tmp_path / "notes.txt"
    path.write_text("Hello from a text file.\nSecond line.", encoding="utf-8")
    return path


@pytest.fixture
def sample_markdown(tmp_path: Path) -> Path:
    path = tmp_path / "readme.md"
    path.write_text("# Title\n\nMarkdown **content**.", encoding="utf-8")
    return path


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    path = tmp_path / "report.docx"
    document = DocxDocument()
    document.add_paragraph("Hello from DOCX.")
    document.add_paragraph("Second paragraph.")
    document.save(path)
    return path
