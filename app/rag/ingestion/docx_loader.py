"""DOCX document loader."""

from pathlib import Path

from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError

from app.rag.ingestion.base import DocumentLoader
from app.rag.ingestion.exceptions import DocumentLoadError
from app.rag.ingestion.models import Document


class DocxLoader(DocumentLoader):
    """Load text content from Microsoft Word DOCX files."""

    @property
    def supported_extensions(self) -> tuple[str, ...]:
        return (".docx",)

    def load(self, source: str | Path) -> Document:
        path = self._resolve_path(source)
        try:
            document = DocxDocument(str(path))
            paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs]
            text = "\n".join(paragraph for paragraph in paragraphs if paragraph)
        except PackageNotFoundError as exc:
            raise DocumentLoadError(f"Invalid DOCX file: {path}") from exc
        except OSError as exc:
            raise DocumentLoadError(f"Failed to open DOCX: {path}") from exc

        return Document(
            text=text,
            metadata=self._base_metadata(
                path,
                format="docx",
                paragraph_count=len(document.paragraphs),
            ),
        )
